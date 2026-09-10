#!/usr/bin/env python3
"""Convert every sheet in the NKI Excel workbook to JSON."""

from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SRC = Path("/Users/umarsaeed/Downloads/2026-09-10_Variables and weights_updated.xlsx")
OUT_DIR = ROOT / "json"

YEARS = [2018, 2019, 2020, 2021, 2022, 2023, 2024, 2025, 2026]


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        if text == "" or text.lower() in {"none", "n/a"}:
            return None
        return text
    if isinstance(value, bool):
        return value
    if isinstance(value, float):
        if value != value:
            return None
        rounded = round(value, 6)
        if rounded == int(rounded) and abs(rounded) < 1e12:
            return int(rounded)
        return rounded
    if isinstance(value, int):
        return value
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return value


def unique_headers(headers):
    seen = {}
    names = []
    for header in headers:
        name = str(header).strip() if header not in (None, "") else "column"
        count = seen.get(name, 0) + 1
        seen[name] = count
        if count == 1:
            names.append(name)
        elif count == 2:
            names.append(f"{name}_scale100")
        else:
            names.append(f"{name}_{count}")
    return names


def write_json(path, payload, pretty=True):
    path.parent.mkdir(parents=True, exist_ok=True)
    if pretty:
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    else:
        with path.open("w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False, separators=(",", ":"))
            handle.write("\n")
    print("wrote", path.name, "mb", round(path.stat().st_size / 1024 / 1024, 2))


def convert_total_effect(ws):
    rows = list(ws.iter_rows(values_only=True))
    note = None
    factors = []
    for row in rows[1:]:
        values = [clean(cell) for cell in row]
        if not values or values[0] is None:
            continue
        item = {"factor": values[0]}
        for year, raw in zip(YEARS, values[1:10]):
            item[str(year)] = clean(raw)
        factors.append(item)
    header = rows[0] if rows else []
    for cell in header:
        if isinstance(cell, str) and "total effect" in cell.lower():
            note = cell.strip()
            break
    return {
        "sheet": "TotalEffect_Factors",
        "note": note,
        "factors": factors,
    }


def convert_weights(ws):
    rows = list(ws.iter_rows(values_only=True))
    note = None
    weights = []
    impact = []
    for row in rows[1:]:
        values = list(row)
        left_key = clean(values[0]) if values else None
        for cell in values:
            if isinstance(cell, str) and "weight" in cell.lower() and "question" in cell.lower():
                note = cell.strip()
        if left_key:
            item = {"question": left_key}
            for year, raw in zip(YEARS, values[1:10]):
                item[str(year)] = clean(raw)
            weights.append(item)
        right_key = clean(values[11]) if len(values) > 11 else None
        if right_key:
            item = {"question": right_key}
            for year, raw in zip(YEARS[:-1], values[12:20]):
                item[str(year)] = clean(raw)
            impact.append(item)
    return {
        "sheet": "Weights_questions",
        "note": note,
        "weights": weights,
        "impactWeights": impact,
    }


def convert_survey(ws, path):
    iterator = ws.iter_rows(values_only=True)
    sections = next(iterator)
    next(iterator)
    headers = unique_headers(next(iterator))
    columns = []
    current_section = None
    for index, name in enumerate(headers):
        section = clean(sections[index]) if index < len(sections) else None
        if section:
            current_section = section
        columns.append({"name": name, "section": current_section})

    path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with path.open("w", encoding="utf-8") as handle:
        handle.write("{\n")
        handle.write(json.dumps("sheet", ensure_ascii=False) + ': "SurveyData 2018-2026",\n')
        handle.write(json.dumps("columns", ensure_ascii=False) + ": ")
        json.dump(columns, handle, ensure_ascii=False)
        handle.write(",\n")
        handle.write('"rows": [\n')
        first = True
        for row in iterator:
            record = {}
            empty = True
            for name, raw in zip(headers, row):
                value = clean(raw)
                if value is None:
                    continue
                empty = False
                record[name] = value
            if empty:
                continue
            if not first:
                handle.write(",\n")
            json.dump(record, handle, ensure_ascii=False, separators=(",", ":"))
            first = False
            count += 1
        handle.write("\n]\n}\n")
    print("wrote", path.name, "mb", round(path.stat().st_size / 1024 / 1024, 2), "rows", count)


def main():
    if not SRC.exists():
        raise SystemExit(f"missing workbook: {SRC}")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SRC, read_only=True, data_only=True)
    print("sheets", wb.sheetnames)

    write_json(OUT_DIR / "TotalEffect_Factors.json", convert_total_effect(wb["TotalEffect_Factors"]))
    write_json(OUT_DIR / "Weights_questions.json", convert_weights(wb["Weights_questions"]))
    convert_survey(wb["SurveyData 2018-2026"], OUT_DIR / "SurveyData_2018-2026.json")
    wb.close()


if __name__ == "__main__":
    main()
